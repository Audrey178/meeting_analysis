"""
docx_writer.py

Renders the generated "Thông báo kết luận" content into a .docx file that
follows the fixed administrative document skeleton:

    THÔNG BÁO
    <trích yếu, in đậm, căn giữa>
    -----
    <đoạn mở>
    I. THÀNH PHẦN THAM DỰ
    ...
    II. NỘI DUNG
    1. <mục>
    2. <mục>
    ...
    <đoạn kết>

The skeleton itself (title, section labels, numbering) is fixed; only the
text passed in for each slot varies per meeting.
"""

import logging
from typing import Dict, List, Tuple

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH

logging.basicConfig(level=logging.INFO, format='%(asctime)s - %(levelname)s - %(message)s')

_NO_PARTICIPANTS_PLACEHOLDER = "[CẦN BỔ SUNG THÀNH PHẦN THAM DỰ]"


def _add_centered_bold(document: Document, text: str, size=None):
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run(text)
    run.bold = True
    if size:
        run.font.size = size
    return paragraph


def _add_heading_bold(document: Document, text: str):
    paragraph = document.add_paragraph()
    run = paragraph.add_run(text)
    run.bold = True
    return paragraph


def write_thong_bao_docx(
    output_path: str,
    title_brief: str,
    opening_paragraph: str,
    participants: Dict,
    content_sections: List[Tuple[str, str]],
    closing_paragraph: str,
) -> str:
    """
    Build and save a "Thông báo kết luận" .docx file.

    Args:
        output_path: path to save the .docx to.
        title_brief: the multi-line "trích yếu" under the THÔNG BÁO title
            (e.g. "Ý kiến kết luận của đồng chí ... tại cuộc họp ..."),
            already fully composed text (may contain '\n').
        opening_paragraph: the opening narrative paragraph
            ("Ngày ... tại ..., đồng chí ... đã chủ trì ...").
        participants: dict with keys:
            - "chair": str — "Đồng chí <tên>, <chức danh>."
            - "attendees": List[str] — bullet lines for "2. Tham dự:"
        content_sections: ordered list of (title, body) tuples — each
            becomes one numbered item under "II. NỘI DUNG".
        closing_paragraph: the closing narrative paragraph.

    Returns:
        The output_path the file was written to.
    """
    document = Document()

    _add_centered_bold(document, "THÔNG BÁO")

    for line in (title_brief or "").split("\n"):
        if line.strip():
            _add_centered_bold(document, line.strip())

    _add_centered_bold(document, "-----")

    document.add_paragraph(opening_paragraph or "[CẦN BỔ SUNG ĐOẠN MỞ ĐẦU]")

    _add_heading_bold(document, "I. THÀNH PHẦN THAM DỰ")

    chair_line = (participants or {}).get("chair") or "[CẦN BỔ SUNG THÔNG TIN CHỦ TRÌ]"
    document.add_paragraph(f"1. Chủ trì: {chair_line}")

    document.add_paragraph("2. Tham dự:")
    attendees = (participants or {}).get("attendees") or [_NO_PARTICIPANTS_PLACEHOLDER]
    for attendee in attendees:
        document.add_paragraph(f"- {attendee}")

    _add_heading_bold(document, "II. NỘI DUNG")

    for index, (title, body) in enumerate(content_sections, start=1):
        header = f"{index}. {title}".strip()
        _add_heading_bold(document, header)
        for line in (body or "[CẦN BỔ SUNG NỘI DUNG]").split("\n"):
            if line.strip():
                document.add_paragraph(line.strip())

    document.add_paragraph(closing_paragraph or "[CẦN BỔ SUNG ĐOẠN KẾT]")

    document.save(output_path)
    logging.info(f"Saved Thông báo kết luận document to {output_path}")
    return output_path
